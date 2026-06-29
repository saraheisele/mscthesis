"""Extract mating-related notes and timestamps from lab session metadata."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document

from data_paths import EXCEL_DIR, LAB_DATA_DIR

# High-confidence mating / courtship vocabulary (English + German).
STRONG_MATING_RE = re.compile(
    r"mating|courtship|spawn|breed|copul|reproduc|fortpflanz|"
    r"paarungsverhalten|paarungsverh\.?|paarung|paaren|pärchen|"
    r"körperkontakt|kuscheln|eiablage|nestbau|"
    r"body\s+contacts?|contact\s+scenes?|tentative\s+contacts?|"
    r"nose\s+to\s+nose|like\s+each\s+other",
    re.IGNORECASE,
)

# Broader terms requested by user — matched only outside negative context.
WEAK_MATING_RE = re.compile(
    r"\bpaar\b|gemeinsam|zusammen|\bkontakt\b|\bcontacts?\b|\bnest\b|\bpair\b",
    re.IGNORECASE,
)

NEGATIVE_CONTEXT_RE = re.compile(
    r"e-?mail|contacted preferably|contact problems|electrode|"
    r"\brepair\b|filter cleaning|not working|eellogger|battery|synchron|"
    r"skin contact male with food|in contact with the water|"
    r"introduced objects|no clear sign of courtship|nothing else seems|"
    r"christmas|happy new year|private email",
    re.IGNORECASE,
)

TIME_RE = re.compile(r"\b(\d{1,2}:\d{2}(?::\d{2})?)\b")
DATE_INLINE_RE = re.compile(r"\b(20\d{6})\b")
SESSION_DATE_RE = re.compile(
    r"recordings_(\d{4}-\d{2}-\d{2})(?:_(\d{2})-(\d{2})-(\d{2}))?"
)


def is_mating_line(line: str) -> bool:
    if not line or not line.strip():
        return False
    has_strong = bool(STRONG_MATING_RE.search(line))
    has_weak = bool(WEAK_MATING_RE.search(line))
    if not has_strong and not has_weak:
        return False
    if re.search(r"no clear sign of courtship", line, re.IGNORECASE):
        return False
    if not has_strong and NEGATIVE_CONTEXT_RE.search(line):
        return False
    return True


def parse_session_date(session_name: str) -> datetime | None:
    match = SESSION_DATE_RE.match(session_name)
    if not match:
        return None
    date_str = match.group(1)
    if match.group(2):
        return datetime.strptime(
            f"{date_str} {match.group(2)}:{match.group(3)}:{match.group(4)}",
            "%Y-%m-%d %H:%M:%S",
        )
    return datetime.strptime(date_str, "%Y-%m-%d")


def parse_time_on_date(time_str: str, base_date: datetime) -> datetime | None:
    parts = [int(x) for x in time_str.split(":")]
    if len(parts) == 2:
        hour, minute = parts
        second = 0
    else:
        hour, minute, second = parts
    if not (0 <= hour <= 23 and 0 <= minute <= 59 and 0 <= second <= 59):
        return None
    return base_date.replace(hour=hour, minute=minute, second=second, microsecond=0)


def _event_times_for_line(line: str, default_date: datetime) -> list[datetime]:
    current_date = default_date.replace(hour=0, minute=0, second=0, microsecond=0)
    inline_date = DATE_INLINE_RE.search(line.replace(" ", ""))
    if inline_date:
        try:
            current_date = datetime.strptime(inline_date.group(1), "%Y%m%d")
        except ValueError:
            pass

    times = TIME_RE.findall(line)
    if times:
        return [
            dt
            for time_str in times
            if (dt := parse_time_on_date(time_str, current_date)) is not None
        ]
    return [default_date.replace(hour=12, minute=0, second=0, microsecond=0)]


def _append_event(
    events: list[dict],
    *,
    session: str,
    source_file: str,
    source_type: str,
    note: str,
    default_date: datetime,
) -> None:
    if not is_mating_line(note):
        return
    for event_time in _event_times_for_line(note, default_date):
        events.append(
            {
                "session": session,
                "source_file": source_file,
                "source_type": source_type,
                "note": note[:300],
                "event_time": event_time,
            }
        )


def _docx_lines(docx_path: Path) -> list[str]:
    doc = Document(docx_path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [line.strip() for line in lines if line.strip()]


def _readme_lines(readme_path: Path) -> list[str]:
    if readme_path.suffix.lower() == ".docx":
        return _docx_lines(readme_path)
    return [
        line.strip()
        for line in readme_path.read_text(encoding="utf-8", errors="ignore").splitlines()
        if line.strip()
    ]


def _excel_default_years(path: Path) -> list[int]:
    """Infer calendar year(s) covered by an environmental log workbook."""
    name = path.stem
    if m := re.search(r"(\d{4})$", name):
        return [int(m.group(1))]
    if m := re.search(r"(\d{2})(\d{2})(\d{2})-(\d{2})(\d{2})(\d{2})", name):
        start_year = 2000 + int(m.group(1))
        end_year = 2000 + int(m.group(4))
        return list(range(start_year, end_year + 1))
    return []


def _parse_excel_date(value, years: list[int], last_year: int | None) -> tuple[datetime | None, int | None]:
    if pd.isna(value):
        return None, last_year

    if isinstance(value, datetime):
        dt = value.replace(hour=12, minute=0, second=0, microsecond=0)
        return dt, dt.year

    text = str(value).strip()
    if not text:
        return None, last_year

    parsed = pd.to_datetime(text, format="mixed", errors="coerce")
    if pd.notna(parsed):
        dt = parsed.to_pydatetime().replace(hour=12, minute=0, second=0, microsecond=0)
        return dt, dt.year

    day_month = re.match(r"^(\d{1,2})\.(\d{1,2})\.?$", text)
    if day_month and years:
        day, month = int(day_month.group(1)), int(day_month.group(2))
        year = last_year or years[0]
        try:
            dt = datetime(year, month, day, 12, 0, 0)
        except ValueError:
            return None, last_year
        if len(years) > 1:
            if last_year is not None and dt.year == last_year and month < 3 and years[-1] > years[0]:
                pass
            elif dt.year == years[0] and month >= 10 and years[-1] > years[0]:
                dt = datetime(years[-1], month, day, 12, 0, 0)
        return dt, dt.year

    return None, last_year


def _excel_remark_events(excel_path: Path) -> list[dict]:
    events = []
    raw = pd.read_excel(excel_path, sheet_name=0, header=None)
    if raw.empty:
        return events

    years = _excel_default_years(excel_path)
    headers = [str(x).lower().strip() if pd.notna(x) else "" for x in raw.iloc[0]]
    date_idx = next(
        (i for i, h in enumerate(headers) if h in {"date", "datum"} or "date" in h or "datum" in h),
        0,
    )
    remark_idx = next(
        (
            i
            for i, h in enumerate(headers)
            if "remark" in h or "anmerkung" in h or "anmerkungen" in h
        ),
        None,
    )

    last_year: int | None = years[0] if years else None
    for row_idx in range(1, len(raw)):
        row = raw.iloc[row_idx]
        default_date, last_year = _parse_excel_date(row[date_idx], years, last_year)
        if default_date is None:
            continue

        if remark_idx is not None and pd.notna(row[remark_idx]):
            note = str(row[remark_idx]).strip()
        else:
            note = " | ".join(
                str(value).strip()
                for col_idx, value in enumerate(row)
                if col_idx != date_idx and pd.notna(value) and str(value).strip()
            )

        if not note:
            continue

        session = f"environment_log_{default_date.strftime('%Y-%m-%d')}"
        _append_event(
            events,
            session=session,
            source_file=excel_path.name,
            source_type="xlsx",
            note=note,
            default_date=default_date,
        )

    return events


def extract_mating_events() -> list[dict]:
    """Return mating-related notes with best-effort event timestamps."""
    events: list[dict] = []

    for session_dir in sorted(LAB_DATA_DIR.glob("recordings_*")):
        images_dir = session_dir / "images"
        if not images_dir.exists():
            continue

        default_date = parse_session_date(session_dir.name)
        if default_date is None:
            continue

        for docx_path in sorted(images_dir.glob("*.docx")):
            if docx_path.name.upper().startswith("README"):
                continue
            try:
                lines = _docx_lines(docx_path)
            except Exception:
                continue
            for line in lines:
                _append_event(
                    events,
                    session=session_dir.name,
                    source_file=docx_path.name,
                    source_type="docx",
                    note=line,
                    default_date=default_date,
                )

        for readme_path in sorted(images_dir.glob("README*")):
            if readme_path.name.endswith("~"):
                continue
            if readme_path.suffix.lower() not in {".md", ".docx"}:
                continue
            try:
                lines = _readme_lines(readme_path)
            except Exception:
                continue
            for line in lines:
                _append_event(
                    events,
                    session=session_dir.name,
                    source_file=readme_path.name,
                    source_type="readme",
                    note=line,
                    default_date=default_date,
                )

    for excel_path in sorted(EXCEL_DIR.glob("*Verlauf-relLeitwert*.xlsx")):
        try:
            events.extend(_excel_remark_events(excel_path))
        except Exception:
            continue

    dedup = {
        (e["session"], e["event_time"], e["source_file"], e["note"]): e
        for e in events
    }
    return sorted(dedup.values(), key=lambda e: e["event_time"])
