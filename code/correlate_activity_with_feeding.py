"""Extract feeding times from session Word records and correlate with pulse activity.

Analysis part: feeding correlation (Part 4 of Berlin activity analysis).
Dependencies: data_paths; requires .h5 files and session .docx logs in LAB_DATA_DIR.

Parses feeding timestamps from session Word documents, aligns them with predetected
pulse recordings, and tests whether pulse rates differ around feeding events.
"""

from __future__ import annotations

import re
import warnings
from datetime import datetime, timedelta
from pathlib import Path

import matplotlib.pyplot as plt
import nixio
import numpy as np
import pandas as pd
from docx import Document
from scipy import stats

warnings.filterwarnings("ignore")

from data_paths import (
    FEEDING_CORRELATION_DIR,
    H5_DIR,
    LAB_DATA_DIR,
)

OUTPUT_DIR = FEEDING_CORRELATION_DIR
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

PULSE_TYPES = {
    "all": {"label": "All pulses", "array": None},
    "double": {"label": "Double pulses", "array": "is_double_peak"},
    "wide": {"label": "Wide pulses", "array": "is_wide_pulse"},
    "fat": {"label": "Fat pulses", "array": "is_fat_pulse"},
}

FEEDING_LINE_RE = re.compile(
    r"(feed|food|strike|füt|fut|pr[äa]sent|present|fress)", re.IGNORECASE
)
NEGATIVE_CONTEXT_RE = re.compile(
    r"(expecting|before feeding|after feeding|not hungry|without feeding|"
    r"introduced objects before|going after any introduced|filter|cleaning|"
    r"eellogger|camera|video|temperature|temperatur|^T\s*=|^L\s*=|rain|"
    r"reference time|recording time|overnight|electrode|battery|synchron)",
    re.IGNORECASE,
)
TIME_RE = re.compile(r"\b(\d{1,2}:\d{2}(?::\d{2})?)\b")
DATE_INLINE_RE = re.compile(r"\b(20\d{6})\b")
SESSION_DATE_RE = re.compile(
    r"recordings_(\d{4}-\d{2}-\d{2})(?:_(\d{2})-(\d{2})-(\d{2}))?"
)
WAV_TIME_RE = re.compile(r"(\d{8})T(\d{6})")

PERI_WINDOW_MIN = 10
FEEDING_FLAG_RADIUS_MIN = 5
BASELINE_EXCLUDE_MIN = 5


def parse_session_date(session_name: str) -> datetime | None:
    match = SESSION_DATE_RE.match(session_name)
    if not match:
        return None
    date_str = match.group(1)
    if match.group(2):
        return datetime.strptime(
            f"{date_str} {match.group(2)}:{match.group(3)}:{match.group(4)}",
            "%Y-%m-%d %H:%M:%S",
        )
    return datetime.strptime(date_str, "%Y-%m-%d")


def parse_time_on_date(time_str: str, base_date: datetime) -> datetime | None:
    parts = [int(x) for x in time_str.split(":")]
    if len(parts) == 2:
        hour, minute = parts
        second = 0
    else:
        hour, minute, second = parts
    if not (0 <= hour <= 23 and 0 <= minute <= 59 and 0 <= second <= 59):
        return None
    return base_date.replace(hour=hour, minute=minute, second=second, microsecond=0)


def docx_text(docx_path: Path) -> str:
    doc = Document(docx_path)
    chunks = [para.text for para in doc.paragraphs if para.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)


def extract_eellogger_on_times(session_name: str, docx_path: Path) -> list[datetime]:
    text = docx_text(docx_path)
    default_date = parse_session_date(session_name)
    if default_date is None:
        return []

    current_date = default_date.replace(hour=0, minute=0, second=0, microsecond=0)
    on_times = []

    for line in text.splitlines():
        inline_date = DATE_INLINE_RE.search(line.replace(" ", ""))
        if inline_date:
            try:
                current_date = datetime.strptime(inline_date.group(1), "%Y%m%d")
            except ValueError:
                pass

        if not re.search(r"e+e?llogger\s+on", line, re.IGNORECASE):
            continue
        for time_str in TIME_RE.findall(line):
            dt = parse_time_on_date(time_str, current_date)
            if dt is not None:
                on_times.append(dt)
    return sorted(set(on_times))


def extract_feeding_events(session_name: str, docx_path: Path) -> list[dict]:
    text = docx_text(docx_path)
    default_date = parse_session_date(session_name)
    if default_date is None:
        return []

    current_date = default_date.replace(hour=0, minute=0, second=0, microsecond=0)
    events = []

    for line in text.splitlines():
        inline_date = DATE_INLINE_RE.search(line.replace(" ", ""))
        if inline_date:
            try:
                current_date = datetime.strptime(inline_date.group(1), "%Y%m%d")
            except ValueError:
                pass

        if not FEEDING_LINE_RE.search(line):
            continue
        if NEGATIVE_CONTEXT_RE.search(line) and not re.search(
            r"(feed|food|strike|fress)", line, re.IGNORECASE
        ):
            continue

        times = TIME_RE.findall(line)
        if not times:
            continue

        event_type = "feeding"
        lower = line.lower()
        if re.search(r"present|pr[äa]sent|provided|entered|entered", lower):
            event_type = "food_presentation"
        elif "strike" in lower:
            event_type = "strike"

        for time_str in times:
            event_dt = parse_time_on_date(time_str, current_date)
            if event_dt is None:
                continue
            events.append(
                {
                    "session": session_name,
                    "docx_file": docx_path.name,
                    "event_type": event_type,
                    "event_time": event_dt,
                    "source_line": line.strip(),
                }
            )

    dedup = {}
    for event in events:
        key = (event["session"], event["event_time"], event["event_type"], event["source_line"])
        dedup[key] = event
    return sorted(dedup.values(), key=lambda e: e["event_time"])


def find_docx_for_session(session_name: str) -> Path | None:
    images_dir = LAB_DATA_DIR / session_name / "images"
    if not images_dir.exists():
        return None
    docx_files = sorted(images_dir.glob("*.docx"))
    return docx_files[0] if docx_files else None


def session_name_from_h5(h5_path: Path) -> str:
    return h5_path.stem.replace("_pulses", "")


def get_wav_times(session_name: str) -> list[datetime]:
    session_dir = LAB_DATA_DIR / session_name
    if not session_dir.exists():
        base_date = parse_session_date(session_name)
        if base_date is None:
            return []
        session_dir = LAB_DATA_DIR / session_name.split("_")[0]
        if not session_dir.exists():
            parent = SESSION_DATE_RE.match(session_name)
            if parent:
                session_dir = LAB_DATA_DIR / f"recordings_{parent.group(1)}"
    if not session_dir.exists():
        return []

    wav_times = []
    for wav_path in session_dir.glob("eellogger*.wav"):
        match = WAV_TIME_RE.search(wav_path.name)
        if match:
            wav_times.append(
                datetime.strptime(match.group(1) + match.group(2), "%Y%m%d%H%M%S")
            )
    return sorted(wav_times)


def infer_recording_start(
    h5_path: Path,
    feeding_events: list[dict] | None = None,
    eellogger_on_times: list[datetime] | None = None,
) -> tuple[datetime, float, str]:
    session_name = session_name_from_h5(h5_path)
    filename_start = parse_session_date(session_name)

    with nixio.File.open(str(h5_path)) as nix_file:
        meta = nix_file.sections["pulses_metadata"]["metadata"]
        duration = float(meta["duration"])
        h5_start = datetime.strptime(
            meta["metadata"]["INFO"]["DateTimeOriginal"], "%Y-%m-%dT%H:%M:%S"
        )

    if filename_start and SESSION_DATE_RE.match(session_name).group(2):
        return filename_start, duration, "h5_filename"

    wav_times = get_wav_times(session_name)
    feeding_events = feeding_events or []
    eellogger_on_times = eellogger_on_times or []

    def score_candidate(candidate: datetime) -> tuple:
        wav_diff = (
            min(abs((w - candidate).total_seconds()) for w in wav_times)
            if wav_times
            else 99999.0
        )
        rec_end = candidate + timedelta(seconds=duration)
        n_feeding = sum(
            1
            for event in feeding_events
            if candidate - timedelta(minutes=FEEDING_FLAG_RADIUS_MIN)
            <= event["event_time"]
            <= rec_end + timedelta(minutes=FEEDING_FLAG_RADIUS_MIN)
        )
        eel_diff = (
            min(abs((t - candidate).total_seconds()) for t in eellogger_on_times)
            if eellogger_on_times
            else 99999.0
        )
        return (wav_diff > 300, -n_feeding, wav_diff, eel_diff)

    candidates = [h5_start, h5_start - timedelta(hours=12), h5_start + timedelta(hours=12)]
    best = min(candidates, key=score_candidate)
    wav_diff = (
        min(abs((w - best).total_seconds()) for w in wav_times) if wav_times else float("inf")
    )
    if wav_diff <= 300:
        return best, duration, "wav_aligned"
    return h5_start, duration, "h5_metadata"


def load_pulses_by_type(
    h5_path: Path,
    feeding_events: list[dict] | None = None,
    eellogger_on_times: list[datetime] | None = None,
) -> tuple[dict[str, np.ndarray], float, datetime, float]:
    with nixio.File.open(str(h5_path)) as nix_file:
        block = nix_file.blocks["pulses"]
        array_names = {da.name for da in block.data_arrays}
        if "centers" not in array_names:
            return {}, 0.0, datetime.min, 0.0

        centers = block.data_arrays["centers"][:]
        if "predicted_labels" in array_names:
            mask = block.data_arrays["predicted_labels"][:] == 1
        else:
            mask = np.ones(len(centers), dtype=bool)

        pulses = {}
        for key, cfg in PULSE_TYPES.items():
            marker = cfg["array"]
            if marker is None:
                pulses[key] = centers[mask]
            elif marker in array_names:
                type_mask = block.data_arrays[marker][:][mask].astype(bool)
                pulses[key] = centers[mask][type_mask]
            else:
                pulses[key] = centers[:0]

        meta = nix_file.sections["pulses_metadata"]["metadata"]
        fs = float(meta["samplerate"])
        duration = float(meta["duration"])

    rec_start, _, _ = infer_recording_start(
        h5_path, feeding_events=feeding_events, eellogger_on_times=eellogger_on_times
    )
    return pulses, fs, rec_start, duration


def minute_pulse_rates(
    pulse_samples: np.ndarray, fs: float, rec_start: datetime, duration: float
) -> pd.DataFrame:
    n_bins = int(np.ceil(duration / 60.0))
    counts = np.zeros(n_bins, dtype=float)
    if len(pulse_samples) > 0:
        seconds = pulse_samples / fs
        bin_idx = (seconds // 60).astype(int)
        bin_idx = bin_idx[(bin_idx >= 0) & (bin_idx < n_bins)]
        if len(bin_idx) > 0:
            counts += np.bincount(bin_idx, minlength=n_bins)[:n_bins]

    timestamps = [rec_start + timedelta(minutes=i) for i in range(n_bins)]
    return pd.DataFrame(
        {
            "timestamp": timestamps,
            "pulse_count": counts,
            "pulse_rate_hz": counts / 60.0,
        }
    )


def events_in_recording(events: list[dict], rec_start: datetime, rec_end: datetime) -> list[dict]:
    matched = []
    for event in events:
        t = event["event_time"]
        if rec_start - timedelta(minutes=FEEDING_FLAG_RADIUS_MIN) <= t <= rec_end + timedelta(
            minutes=FEEDING_FLAG_RADIUS_MIN
        ):
            matched.append(event)
    return matched


def add_feeding_flags(minute_df: pd.DataFrame, events: list[dict]) -> pd.DataFrame:
    out = minute_df.copy()
    out["feeding_flag"] = 0
    out["minutes_to_feeding"] = np.nan
    if not events:
        return out

    for idx, row in out.iterrows():
        diffs = [
            (event["event_time"] - row["timestamp"]).total_seconds() / 60.0
            for event in events
        ]
        nearest = min(diffs, key=abs)
        out.at[idx, "minutes_to_feeding"] = nearest
        if abs(nearest) <= FEEDING_FLAG_RADIUS_MIN:
            out.at[idx, "feeding_flag"] = 1
    return out


def peri_event_trajectory(
    pulse_samples: np.ndarray, fs: float, rec_start: datetime, events: list[dict]
) -> np.ndarray | None:
    if len(events) == 0 or len(pulse_samples) == 0:
        return None

    window_s = PERI_WINDOW_MIN * 60
    bin_edges = np.arange(-window_s, window_s + 60, 60)
    trajectories = []

    for event in events:
        offset_s = (event["event_time"] - rec_start).total_seconds()
        rel_seconds = pulse_samples / fs - offset_s
        hist, _ = np.histogram(rel_seconds, bins=bin_edges)
        trajectories.append(hist / 60.0)

    if not trajectories:
        return None
    return np.mean(np.vstack(trajectories), axis=0)


def correlate_feeding_flag(minute_records: pd.DataFrame) -> dict:
    if minute_records.empty or minute_records["feeding_flag"].nunique() < 2:
        return {
            "n_minutes": len(minute_records),
            "n_feeding_minutes": int(minute_records["feeding_flag"].sum()),
            "pearson_r": np.nan,
            "pearson_p": np.nan,
            "spearman_r": np.nan,
            "spearman_p": np.nan,
            "mean_rate_feeding_hz": np.nan,
            "mean_rate_nonfeeding_hz": np.nan,
            "rate_ratio_feeding_over_nonfeeding": np.nan,
        }

    feeding = minute_records["feeding_flag"].astype(float)
    rate = minute_records["pulse_rate_hz"].astype(float)
    pearson_r, pearson_p = stats.pearsonr(feeding, rate)
    spearman_r, spearman_p = stats.spearmanr(feeding, rate)

    feeding_rates = minute_records.loc[minute_records["feeding_flag"] == 1, "pulse_rate_hz"]
    nonfeeding_rates = minute_records.loc[
        (minute_records["feeding_flag"] == 0)
        & (
            minute_records["minutes_to_feeding"].isna()
            | (minute_records["minutes_to_feeding"].abs() > BASELINE_EXCLUDE_MIN)
        ),
        "pulse_rate_hz",
    ]

    mean_feed = feeding_rates.mean() if len(feeding_rates) else np.nan
    mean_non = nonfeeding_rates.mean() if len(nonfeeding_rates) else np.nan
    ratio = mean_feed / mean_non if mean_non and not np.isnan(mean_non) else np.nan

    return {
        "n_minutes": len(minute_records),
        "n_feeding_minutes": int(minute_records["feeding_flag"].sum()),
        "pearson_r": pearson_r,
        "pearson_p": pearson_p,
        "spearman_r": spearman_r,
        "spearman_p": spearman_p,
        "mean_rate_feeding_hz": mean_feed,
        "mean_rate_nonfeeding_hz": mean_non,
        "rate_ratio_feeding_over_nonfeeding": ratio,
    }


def extract_all_feeding_events() -> pd.DataFrame:
    rows = []
    for session_dir in sorted(LAB_DATA_DIR.glob("recordings_*")):
        if not session_dir.is_dir():
            continue
        docx_path = find_docx_for_session(session_dir.name)
        if docx_path is None:
            continue
        for event in extract_feeding_events(session_dir.name, docx_path):
            rows.append(event)

    if not rows:
        return pd.DataFrame()

    df = pd.DataFrame(rows)
    df["event_time"] = pd.to_datetime(df["event_time"])
    return df.sort_values(["session", "event_time"]).reset_index(drop=True)


def run_analysis():
    print("Extracting feeding times from Word documents...")
    feeding_df = extract_all_feeding_events()
    feeding_df.to_csv(OUTPUT_DIR / "feeding_events_extracted.csv", index=False)
    write_session_feeding_summary(feeding_df)
    print(f"Extracted {len(feeding_df)} feeding-related events from {feeding_df['session'].nunique()} sessions")

    session_events = {
        session: group.to_dict("records")
        for session, group in feeding_df.groupby("session")
    }
    session_eellogger_on = {}
    for session_dir in sorted(LAB_DATA_DIR.glob("recordings_*")):
        docx_path = find_docx_for_session(session_dir.name)
        if docx_path is None:
            continue
        session_eellogger_on[session_dir.name] = extract_eellogger_on_times(
            session_dir.name, docx_path
        )

    h5_files = sorted(H5_DIR.glob("*_pulses.h5"))
    print(f"Processing {len(h5_files)} h5 files...")

    minute_rows = []
    h5_summary_rows = []
    peri_curves = {key: [] for key in PULSE_TYPES}

    for h5_path in h5_files:
        session_name = session_name_from_h5(h5_path)
        events = session_events.get(session_name, [])
        eellogger_on = session_eellogger_on.get(session_name, [])
        rec_start, duration, start_method = infer_recording_start(
            h5_path, feeding_events=events, eellogger_on_times=eellogger_on
        )
        rec_end = rec_start + timedelta(seconds=duration)
        matched_events = events_in_recording(events, rec_start, rec_end)

        pulses_by_type, fs, _, _ = load_pulses_by_type(
            h5_path, feeding_events=events, eellogger_on_times=eellogger_on
        )
        if not pulses_by_type or "all" not in pulses_by_type:
            continue

        for pulse_type, pulse_samples in pulses_by_type.items():
            minute_df = minute_pulse_rates(pulse_samples, fs, rec_start, duration)
            minute_df = add_feeding_flags(minute_df, matched_events)
            minute_df["h5_file"] = h5_path.name
            minute_df["session"] = session_name
            minute_df["pulse_type"] = pulse_type
            minute_df["recording_start"] = rec_start
            minute_df["recording_end"] = rec_end
            minute_df["n_feeding_events_in_recording"] = len(matched_events)
            minute_rows.append(minute_df)

            if pulse_type == "all":
                corr = correlate_feeding_flag(minute_df)
                h5_summary_rows.append(
                    {
                        "h5_file": h5_path.name,
                        "session": session_name,
                        "recording_start": rec_start,
                        "recording_end": rec_end,
                        "start_time_method": start_method,
                        "n_feeding_events_in_recording": len(matched_events),
                        **corr,
                    }
                )

            curve = peri_event_trajectory(pulse_samples, fs, rec_start, matched_events)
            if curve is not None:
                peri_curves[pulse_type].append(curve)

    minute_records = pd.concat(minute_rows, ignore_index=True) if minute_rows else pd.DataFrame()
    minute_records.to_csv(OUTPUT_DIR / "minute_pulse_rates_with_feeding_flags.csv", index=False)

    h5_summary = pd.DataFrame(h5_summary_rows)
    h5_summary.to_csv(OUTPUT_DIR / "h5_feeding_correlation_summary.csv", index=False)

    corr_rows = []
    for pulse_type in PULSE_TYPES:
        subset = minute_records[minute_records["pulse_type"] == pulse_type]
        corr = correlate_feeding_flag(subset)
        corr_rows.append({"pulse_type": pulse_type, **corr})
    corr_summary = pd.DataFrame(corr_rows)
    corr_summary.to_csv(OUTPUT_DIR / "global_feeding_correlation_by_pulse_type.csv", index=False)

    matched_events_df = []
    for h5_path in h5_files:
        session_name = session_name_from_h5(h5_path)
        events = session_events.get(session_name, [])
        eellogger_on = session_eellogger_on.get(session_name, [])
        rec_start, duration, start_method = infer_recording_start(
            h5_path, feeding_events=events, eellogger_on_times=eellogger_on
        )
        rec_end = rec_start + timedelta(seconds=duration)
        events = session_events.get(session_name, [])
        for event in events_in_recording(events, rec_start, rec_end):
            matched_events_df.append(
                {
                    **event,
                    "h5_file": h5_path.name,
                    "recording_start": rec_start,
                    "recording_end": rec_end,
                    "seconds_from_recording_start": (
                        event["event_time"] - rec_start
                    ).total_seconds(),
                    "start_time_method": start_method,
                }
            )
    pd.DataFrame(matched_events_df).to_csv(
        OUTPUT_DIR / "feeding_events_matched_to_h5.csv", index=False
    )
    write_coverage_report(feeding_df, matched_events_df, h5_summary)

    plot_peri_feeding_curves(peri_curves)
    plot_feeding_vs_nonfeeding_rates(corr_summary)
    print_summary(feeding_df, h5_summary, corr_summary)
    print(f"Results saved to {OUTPUT_DIR}")


def write_session_feeding_summary(feeding_df: pd.DataFrame):
    if feeding_df.empty:
        return

    summary_rows = []
    for session, group in feeding_df.groupby("session"):
        times = "; ".join(
            f"{row.event_time.strftime('%Y-%m-%d %H:%M:%S')} ({row.event_type})"
            for row in group.itertuples()
        )
        summary_rows.append(
            {
                "session": session,
                "docx_file": group.iloc[0]["docx_file"],
                "n_feeding_events": len(group),
                "feeding_times": times,
            }
        )
    pd.DataFrame(summary_rows).sort_values("session").to_csv(
        OUTPUT_DIR / "feeding_times_by_session.csv", index=False
    )


def write_coverage_report(
    feeding_df: pd.DataFrame, matched_events: list[dict], h5_summary: pd.DataFrame
):
    matched_df = pd.DataFrame(matched_events)
    rows = []
    for session in sorted(feeding_df["session"].unique()):
        n_events = int((feeding_df["session"] == session).sum())
        n_matched = (
            int((matched_df["session"] == session).sum()) if not matched_df.empty else 0
        )
        rows.append(
            {
                "session": session,
                "n_feeding_events_in_docx": n_events,
                "n_feeding_events_in_h5_window": n_matched,
                "n_h5_files_for_session": len(list(H5_DIR.glob(f"{session}*_pulses.h5"))),
                "feeding_covered_by_h5": n_matched > 0,
            }
        )
    pd.DataFrame(rows).to_csv(OUTPUT_DIR / "feeding_h5_coverage_report.csv", index=False)


def plot_peri_feeding_curves(peri_curves: dict[str, list[np.ndarray]]):
    if not any(peri_curves.values()):
        return

    fig, ax = plt.subplots(figsize=(10, 6))
    x = np.arange(-PERI_WINDOW_MIN, PERI_WINDOW_MIN)

    for pulse_type, cfg in PULSE_TYPES.items():
        curves = peri_curves[pulse_type]
        if not curves:
            continue
        mean_curve = np.mean(np.vstack(curves), axis=0)
        sem = stats.sem(np.vstack(curves), axis=0) if len(curves) > 1 else np.zeros_like(mean_curve)
        ax.plot(x, mean_curve, label=cfg["label"])
        ax.fill_between(x, mean_curve - sem, mean_curve + sem, alpha=0.2)

    ax.axvline(0, color="black", linestyle="--", linewidth=1, alpha=0.7)
    ax.set_xlabel("Minutes relative to feeding event")
    ax.set_ylabel("Mean pulse rate (Hz)")
    ax.set_title("Average pulse activity around feeding events")
    ax.grid(True, alpha=0.3)
    ax.legend()
    plt.tight_layout()
    plt.savefig(OUTPUT_DIR / "peri_feeding_pulse_rate_trajectories.png", dpi=300)
    plt.close()


def plot_feeding_vs_nonfeeding_rates(corr_summary: pd.DataFrame):
    if corr_summary.empty:
        return

    labels = [PULSE_TYPES[row["pulse_type"]]["label"] for _, row in corr_summary.iterrows()]
    feeding = corr_summary["mean_rate_feeding_hz"].values
    nonfeeding = corr_summary["mean_rate_nonfeeding_hz"].values

    x = np.arange(len(labels))
    width = 0.35
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.bar(x - width / 2, nonfeeding, width, label=f"Non-feeding (>{BASELINE_EXCLUDE_MIN} min away)")
    ax.bar(x + width / 2, feeding, width, label=f"Feeding (±{FEEDING_FLAG_RADIUS_MIN} min)")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=15)
    ax.set_ylabel("Mean pulse rate (Hz)")
    ax.set_title("Pulse rate during feeding windows vs baseline")
    ax.legend()
    ax.grid(True, axis="y", alpha=0.3)
    plt.tight_layout()
    plt.savefig(OUTPUT_DIR / "feeding_vs_nonfeeding_pulse_rates.png", dpi=300)
    plt.close()


def print_summary(feeding_df: pd.DataFrame, h5_summary: pd.DataFrame, corr_summary: pd.DataFrame):
    print("\n=== Feeding extraction summary ===")
    print(f"Sessions with docx records: {feeding_df['session'].nunique() if not feeding_df.empty else 0}")
    print(f"Total extracted events: {len(feeding_df)}")
    if not feeding_df.empty:
        print(feeding_df["event_type"].value_counts().to_string())

    print("\n=== H5 overlap summary ===")
    if not h5_summary.empty:
        with_events = h5_summary[h5_summary["n_feeding_events_in_recording"] > 0]
        print(f"H5 files with feeding events in recording window: {len(with_events)} / {len(h5_summary)}")
        print(
            f"Total feeding events matched to a recording: "
            f"{with_events['n_feeding_events_in_recording'].sum():.0f}"
        )

    print("\n=== Global feeding correlation (minute-level feeding flag vs pulse rate) ===")
    if not corr_summary.empty:
        for _, row in corr_summary.iterrows():
            label = PULSE_TYPES[row["pulse_type"]]["label"]
            print(
                f"{label:16s}  r={row['pearson_r']:+.3f} (p={row['pearson_p']:.4g}), "
                f"feeding={row['mean_rate_feeding_hz']:.3f} Hz, "
                f"non-feeding={row['mean_rate_nonfeeding_hz']:.3f} Hz, "
                f"ratio={row['rate_ratio_feeding_over_nonfeeding']:.2f}"
            )


if __name__ == "__main__":
    run_analysis()
